"""Parser for the ECUSTFD ``food_info.xls`` file.

The original Excel file is stored on disk with the ``.xls`` extension
but is actually an Office Open XML (xlsx) workbook. ``xlrd >= 2.0``
refuses to read it, and ``openpyxl`` can open it directly as a BytesIO
stream without any disk write (avoiding temp-file failures on full disks).

Why a separate module:
- The file is referenced by the runtime MATLAB baseline (``faster_
  rcnn_rec.m``). Extracting its parsing into a dedicated module makes
  it easy to:
  * unit-test the parsing logic without spinning up the rest of the
    pipeline;
  * swap to a hard-coded JSON copy if the .xls file is removed.
"""

from __future__ import annotations

import io # Lưu trữ dữ liệu tạm thời
import zipfile # Đọc file dưới dạng zip
from pathlib import Path # Sử lý đường dẫn file
from typing import Dict, Optional # Sử lý kiểu dữ liệu
    
import openpyxl # Đọc file excel


EXPECTED_HEADER = ["food_name", "shape", "calorie/volume(kC/cm^3)"] # Header của file excel

# Mở file excel dưới dạng xlsx
def _open_as_xlsx(path: Path):
    """Open ``path`` as a real xlsx file even if its extension is .xls.

    The on-disk file has a ZIP header (i.e. it is xlsx) but its
    extension is .xls. We read it into a BytesIO buffer so that openpyxl
    can load it without writing to the filesystem — no temp files, no
    shutil.copy, no disk-space failure.
    """
    # Nếu file có đuôi .xlsx thì mở trực tiếp
    if path.suffix.lower() == ".xlsx":
        return openpyxl.load_workbook(path, data_only=True, read_only=True)# data_only=True để lấy giá trị, read_only=True để đọc file

    # Đọc file dưới dạng zip
    raw_bytes = path.read_bytes() # Đọc file dưới dạng bytes
    buffer = io.BytesIO(raw_bytes) # Đọc file dưới dạng BytesIO

    # Kiểm tra xem có phải là file zip không
    if zipfile.is_zipfile(buffer):
        buffer.seek(0) # Đưa con trỏ về đầu file
        return openpyxl.load_workbook(buffer, data_only=True, read_only=True)# data_only=True để lấy giá trị, read_only=True để đọc file
    else:
        # Nếu không phải file zip thì raise ValueError
        raise ValueError(
            f"{path} does not look like a valid xlsx file (no ZIP header). "
            "Cannot parse as food_info."
        )

# Đọc file excel dưới dạng dict
def parse_food_info(path: str | Path) -> Dict[str, Dict[str, float | str]]: # Trả về dict
    """Parse ``food_info.xls`` into a dict per class.

    Returns:
        ``{"apple": {"shape": "ellipsoid", "kcal_per_cm3": 0.4071}, ...}``.
        The ``coin`` row is included but its values are ``None``.
    """
    p = Path(path) # Path là class để xử lý đường dẫn file
    # Kiểm tra file có tồn tại không
    if not p.exists(): 
        raise FileNotFoundError(p)

    wb = _open_as_xlsx(p) # Mở file excel
    first = wb.sheetnames[0] # Lấy sheet đầu tiên
    sh = wb[first] # Lấy sheet

    rows = list(sh.iter_rows(values_only=True)) # Đọc file excel
    if not rows: # Kiểm tra file có rỗng không
        raise ValueError(f"{p} is empty") # Nếu file rỗng thì raise ValueError

    header = list(rows[0]) # Header của file excel
    if header[:3] != EXPECTED_HEADER: # Kiểm tra header có đúng không
        raise ValueError(
            f"Unexpected header in {p}: {header[:3]} (expected {EXPECTED_HEADER})"
        )

    out: Dict[str, Dict[str, float | str]] = {} # Dict để lưu kết quả
    for row in rows[1:]:
        if not row or row[0] is None: # Kiểm tra hàng có rỗng không
            continue
        name = str(row[0]).strip() # Lấy tên
        shape = None if row[1] is None else str(row[1]).strip() # Lấy hình dạng
        kcal = None if row[2] is None else float(row[2]) # Lấy năng lượng
        out[name] = {"shape": shape, "kcal_per_cm3": kcal} # Lưu kết quả
        #Also expose the snake_case alias used by ``src/constants.py``
        # so callers can look up by either spelling
        # (e.g. "fired_dough_twist" vs "fired dough twist").
        if " " in name: # Kiểm tra tên có chứa khoảng trắng không
            out[name.replace(" ", "_")] = {"shape": shape, "kcal_per_cm3": kcal}

    return out

# Lấy q factor
def get_q_factor(
    food_info: Dict[str, Dict[str, float | str]], class_name: str
) -> Optional[float]: # Trả về float
    """Return the ``kcal_per_cm3`` factor for the class, or ``None``."""
    entry = food_info.get(class_name) # Lấy q factor
    if entry is None: # Kiểm tra q factor có tồn tại không
        return None # Trả về None nếu không tồn tại
    return entry.get("kcal_per_cm3")  # type: ignore[return-value]


__all__ = ["parse_food_info", "get_q_factor", "EXPECTED_HEADER"] # Tất cả các hàm được export
